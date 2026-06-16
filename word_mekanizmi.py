import docx

# Qiymətləndirmə şkalası
def baldan_qiymet_hesabla(bal):
    try:
        # Balı ədədə çeviririk (vergülü nöqtəyə çevirib hesablayırıq)
        b = float(str(bal).replace(',', '.'))
        
        if b >= 71: return "5"
        elif b >= 51: return "4"
        elif b >= 31: return "3"
        else: return "2"
    except:
        return "2" # Əgər bal oxunmazsa, avtomatik 2 yazır

def word_faylini_doldur(word_fayli, axtarilan_fenn, ballar_siyahisi):
    doc = docx.Document(word_fayli)
    table = doc.tables[0]
    
    hedef_sutun_indeksi = -1
    baslangic_setir_indeksi = -1
    
    # Fənnin adını və sütununu tapırıq
    for setir_idx in range(min(5, len(table.rows))):
        for sutun_idx, xana in enumerate(table.rows[setir_idx].cells):
            if axtarilan_fenn.lower() in xana.text.lower():
                hedef_sutun_indeksi = sutun_idx
                baslangic_setir_indeksi = setir_idx + 2 
                break
        if hedef_sutun_indeksi != -1:
            break
            
    if hedef_sutun_indeksi == -1:
        raise ValueError(f"Word cədvəlində '{axtarilan_fenn}' başlığı tapılmadı.")
        
    # İndi həm B (Bal) həm də Q (Qiymət) sütunlarını doldururuq
    bal_idx = 0
    for setir_idx in range(baslangic_setir_indeksi, len(table.rows)):
        if bal_idx < len(ballar_siyahisi):
            bal = ballar_siyahisi[bal_idx]
            qiymet = baldan_qiymet_hesabla(bal)
            
            # B sütununa balı yazırıq (nöqtəni vergüllə əvəz edərək)
            table.cell(setir_idx, hedef_sutun_indeksi).text = str(bal).replace('.', ',')
            
            # Q sütununa (sağdakı xana) qiyməti yazırıq
            if hedef_sutun_indeksi + 1 < len(table.rows[setir_idx].cells):
                table.cell(setir_idx, hedef_sutun_indeksi + 1).text = str(qiymet)
                
            bal_idx += 1
            
    return doc
